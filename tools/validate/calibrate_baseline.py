"""Measures where LibreOffice puts text baselines in a pptx text box
(plan §5.2), to derive the box-top formula the exporter uses.

    uv run --project tools/validate python -I tools/validate/calibrate_baseline.py OUT_DIR

Builds one slide per (font, size, line spacing) with a two-line text box
(zero insets, no wrap, no autofit) whose top is at TOP pt, renders the deck
with LibreOffice and prints, per case, the first baseline below the box top
and the distance between the two baselines, both in em, next to the font's
vertical metrics.

Each slide also carries thin red guide lines where the measured model puts
the two baselines and a caption naming the case, so the same file checks
PowerPoint by eye (Windows milestone check, docs/milestones/).

Fonts Typst embeds (Libertinus Serif, New Computer Modern) are made visible
to LibreOffice through a private fontconfig file; pass their directory with
--extra-fonts (typst-assets' files/fonts in the cargo registry).
"""

import argparse
import ctypes
import os
import pathlib
import subprocess
import tempfile

import pypdfium2 as pdfium
import pypdfium2.raw as raw
from fontTools.ttLib import TTFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Pt

TOP = 100.0
SIZES = [12, 24, 48]
# None = single spacing; numbers are exact spacing as a multiple of the size.
SPACINGS = [None, 0.8, 1.0, 1.2, 1.6]

FONTS = {
    "IBM Plex Sans": "fixtures/fonts/ibm-plex/IBMPlexSans-Regular.ttf",
    "Libertinus Serif": "{extra}/LibertinusSerif-Regular.otf",
    "New Computer Modern": "{extra}/NewCM10-Regular.otf",
    "DejaVu Sans Mono": "{extra}/DejaVuSansMono.ttf",
}


def metrics(path: str) -> dict[str, float]:
    font = TTFont(path)
    upm = font["head"].unitsPerEm
    hhea, os2 = font["hhea"], font["OS/2"]
    return {
        "typo_asc": os2.sTypoAscender / upm,
        "typo_desc": -os2.sTypoDescender / upm,
        "typo_gap": os2.sTypoLineGap / upm,
        "hhea_asc": hhea.ascent / upm,
        "hhea_desc": -hhea.descent / upm,
        "hhea_gap": hhea.lineGap / upm,
        "win_asc": os2.usWinAscent / upm,
        "win_desc": os2.usWinDescent / upm,
        "use_typo": bool(os2.fsSelection & 128),
    }


def build(cases: list[tuple[str, int, float | None]], path: pathlib.Path) -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(841.89), Pt(473.56)
    blank = prs.slide_layouts[6]
    for family, size, spacing in cases:
        slide = prs.slides.add_slide(blank)
        box = slide.shapes.add_textbox(Pt(50), Pt(TOP), Pt(600), Pt(size * 4))
        frame = box.text_frame
        frame.word_wrap = False
        frame.auto_size = MSO_AUTO_SIZE.NONE
        frame.margin_left = frame.margin_right = 0
        frame.margin_top = frame.margin_bottom = 0
        paragraph = frame.paragraphs[0]
        paragraph.space_before = paragraph.space_after = Pt(0)
        if spacing is not None:
            paragraph.line_spacing = Pt(size * spacing)
        run = paragraph.add_run()
        run.text = "Hxg"
        paragraph.add_line_break()
        second = paragraph.add_run()
        second.text = "Hxg"
        for r in (run, second):
            r.font.name = family
            r.font.size = Pt(size)

        # Model (pptx-ir/src/text.rs): first baseline = S - 0.2 em, pitch = S.
        pitch = size * (spacing if spacing is not None else 1.2)
        first = TOP + pitch - 0.2 * size
        for y in (first, first + pitch):
            line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(40), Pt(y), Pt(700), Pt(y))
            line.line.color.rgb = RGBColor(0xE0, 0x20, 0x20)
            line.line.width = Pt(0.5)
        caption = slide.shapes.add_textbox(Pt(50), Pt(30), Pt(600), Pt(20)).text_frame
        label = "single" if spacing is None else f"exactly {spacing:g} em"
        caption.text = f"{family}, {size} pt, line spacing {label}"
        caption.paragraphs[0].runs[0].font.size = Pt(12)
    prs.save(path)


def baselines(pdf_path: pathlib.Path) -> list[tuple[float, float]]:
    """(first, second) baseline y of each page, pt from the page top."""
    out = []
    for page in pdfium.PdfDocument(str(pdf_path)):
        text = page.get_textpage()
        height = page.get_height()
        ys = []
        for i in range(text.count_chars()):
            if raw.FPDFText_GetUnicode(text, i) == ord("H"):
                x, y = ctypes.c_double(), ctypes.c_double()
                raw.FPDFText_GetCharOrigin(text, i, x, y)
                ys.append(height - y.value)
        out.append((ys[0], ys[1]))
    return out


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("out", type=pathlib.Path)
    parser.add_argument("--extra-fonts", required=True, type=pathlib.Path)
    args = parser.parse_args()
    args.out.mkdir(parents=True, exist_ok=True)

    fonts = {name: path.format(extra=args.extra_fonts) for name, path in FONTS.items()}
    cases = [(f, s, sp) for f in fonts for s in SIZES for sp in SPACINGS]
    deck = args.out / "calibration.pptx"
    build(cases, deck)

    with tempfile.TemporaryDirectory() as tmp:
        conf = pathlib.Path(tmp) / "fonts.conf"
        conf.write_text(
            '<?xml version="1.0"?><!DOCTYPE fontconfig SYSTEM "urn:fontconfig:fonts.dtd">'
            "<fontconfig><include ignore_missing=\"yes\">/etc/fonts/fonts.conf</include>"
            f"<dir>{args.extra_fonts.resolve()}</dir></fontconfig>"
        )
        env = dict(os.environ, FONTCONFIG_FILE=str(conf))
        subprocess.run(
            ["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(args.out), str(deck)],
            check=True,
            env=env,
            capture_output=True,
        )

    measured = baselines(args.out / "calibration.pdf")
    print(f"{'font':20} {'size':>4} {'spacing':>7} {'first/em':>8} {'pitch/em':>8}")
    for (family, size, spacing), (b1, b2) in zip(cases, measured):
        label = "single" if spacing is None else f"{spacing:.1f}em"
        print(f"{family:20} {size:4} {label:>7} {(b1 - TOP) / size:8.4f} {(b2 - b1) / size:8.4f}")
    print()
    for family, path in fonts.items():
        m = metrics(path)
        print(family, " ".join(f"{k}={v:.3f}" if isinstance(v, float) else f"{k}={v}" for k, v in m.items()))


if __name__ == "__main__":
    main()
